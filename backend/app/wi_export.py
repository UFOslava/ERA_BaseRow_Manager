import os
import re
import json
import io
import requests
import logging
from datetime import datetime
from docx import Document
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Cm

logger = logging.getLogger(__name__)

from jinja2 import Undefined

class SilentUndefined(Undefined):
    def __getattr__(self, name):
        return self
    def __str__(self):
        return ""
    def __html__(self):
        return ""

class ItemProxy:
    def __init__(self, item_dict, step_number):
        self._item_dict = item_dict
        self._step_number = step_number
        
    def __getattr__(self, name):
        return getattr(self._item_dict, name, self._item_dict.get(name, ""))
        
    def __getitem__(self, key):
        return self._item_dict.get(key, "")
        
    def get(self, key, default=None):
        return self._item_dict.get(key, default)

class ProxyListIterator:
    def __init__(self, flat_items, step_proxy):
        self.flat_items = flat_items
        self.step_proxy = step_proxy
        self.index = 0
        
    def __iter__(self):
        return self
        
    def __next__(self):
        if self.index >= len(self.flat_items):
            raise StopIteration
        item_proxy = self.flat_items[self.index]
        self.step_proxy._active_step_number = item_proxy._step_number
        self.index += 1
        return item_proxy

class ProxyList(list):
    def __init__(self, flat_items, step_proxy):
        super().__init__(flat_items)
        self.flat_items = flat_items
        self.step_proxy = step_proxy
        
    def __iter__(self):
        return ProxyListIterator(self.flat_items, self.step_proxy)

class GlobalStepProxy:
    def __init__(self, context_steps):
        self.context_steps = context_steps
        self._active_step_number = ""
        
    @property
    def tools(self):
        flat_tools = []
        for s in self.context_steps:
            s_num = s.get("step_number", "")
            for t in s.get("tools", []):
                flat_tools.append(ItemProxy(t, s_num))
        return ProxyList(flat_tools, self)
        
    @property
    def parts(self):
        flat_parts = []
        for s in self.context_steps:
            s_num = s.get("step_number", "")
            for p in s.get("parts", []):
                flat_parts.append(ItemProxy(p, s_num))
        return ProxyList(flat_parts, self)
        
    @property
    def step_number(self):
        return self._active_step_number

def get_real_image_url(url, api_url):
    if not url:
        return ""
    if url.startswith("/"):
        from urllib.parse import urlparse, urlunparse
        parsed_api = urlparse(api_url)
        return urlunparse((parsed_api.scheme, parsed_api.netloc, url, "", "", ""))
    if "localhost" in url or "127.0.0.1" in url:
        from urllib.parse import urlparse, urlunparse
        parsed_url = urlparse(url)
        parsed_api = urlparse(api_url)
        return urlunparse((parsed_api.scheme, parsed_api.netloc, parsed_url.path, parsed_url.params, parsed_url.query, parsed_url.fragment))
    return url

KNOWN_TOKENS = {
    "item_pn", "pn", "revision", "description", "ext_pn", "date",
    "step_number", "main_action", "instruction_text",
    "step_image", "item_image",
    "steps", "step", "part", "part_name", "part_qty",
    "tool", "tool_name", "tool_qty", "unique_tools",
    "parts", "tools",
    "for", "endfor", "if", "endif",
    "bom_items", "child"
}

def scan_template(file_path):
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + "\n"
    for section in doc.sections:
        for para in section.header.paragraphs:
            text += para.text + "\n"
        for para in section.footer.paragraphs:
            text += para.text + "\n"
                
    # Find all {{ ... }} and {% ... %}
    pattern = re.compile(r'\{\{(.*?)\}\}|\{\%(.*?)\%\}')
    matches = pattern.findall(text)
    
    found_tokens = set()
    invalid_tokens = set()
    
    for match in matches:
        token = match[0] if match[0] else match[1]
        token = token.strip()
        if not token:
            continue
            
        # Strip pipe filters (support both | width:8cm and | width('8cm') or others)
        if "|" in token:
            token = token.split("|")[0].strip()
            
        # Handle loops
        parts = token.split()
        if parts:
            if parts[0] in ("for", "endfor", "if", "endif"):
                found_tokens.add(parts[0])
                if parts[0] == "for" and len(parts) >= 4:
                    # e.g., for step in steps
                    found_tokens.add(parts[1])
                    found_tokens.add(parts[3])
                    # Strip dots from collection name
                    coll = parts[3].split(".")[0]
                    if coll not in KNOWN_TOKENS and not coll.startswith("step") and not coll.startswith("part") and not coll.startswith("tool"):
                        invalid_tokens.add(parts[3])
                continue
                
        # Basic variable
        base_token = parts[0] if parts else token
        # Handle dot notations
        if "." in base_token:
            parts_dot = base_token.split(".")
            prefix = parts_dot[0]
            prop = parts_dot[1] if len(parts_dot) > 1 else ""
            
            if prefix == "loop":
                continue
            elif prefix == "step":
                found_tokens.add(prefix)
                if prop not in ("step_number", "main_action", "instruction_text", "step_image", "parts", "tools"):
                    invalid_tokens.add(base_token)
                continue
            elif prefix in ("part", "tool", "child"):
                found_tokens.add(prefix)
                if prop not in ("id", "part_number", "pn", "revision", "description", "ext_pn", "quantity", "image"):
                    invalid_tokens.add(base_token)
                continue
            base_token = prefix
            
        found_tokens.add(base_token)
        if base_token not in KNOWN_TOKENS:
            invalid_tokens.add(base_token)
            
    return {
        "valid": len(invalid_tokens) == 0,
        "found": list(found_tokens),
        "invalid": list(invalid_tokens)
    }

def evaluate_instruction_text(step):
    template = step.get("description", "")
    part_slots = step.get("part_slots", [])
    tool_slots = step.get("tool_slots", [])
    action = step.get("action", "")
    receiving_item = step.get("receiving_item")
    receiving_name = receiving_item.get("description") if receiving_item else ""

    def format_slot(slot):
        if not slot or slot.get("id") is None:
            return None
        qty = slot.get("quantity") or 1
        prefix = f"{qty}x " if qty > 1 else ""
        desc = slot.get("description") or "No description"
        pn = slot.get("pn") or slot.get("part_number")
        if pn:
            return f"{prefix}\"{desc}\" ({pn})"
        return f"{prefix}{desc}"

    filled_parts = [format_slot(p) for p in part_slots]
    filled_parts = [p for p in filled_parts if p is not None]
    
    filled_tools = [format_slot(t) for t in tool_slots]
    filled_tools = [t for t in filled_tools if t is not None]

    default_part = ", ".join(filled_parts) if filled_parts else "[Action Item A]"
    default_tool = ", ".join(filled_tools) if filled_tools else "[Tool]"

    text = template
    if receiving_name:
        text = text.replace("{receiving_item}", receiving_name).replace("{b}", receiving_name)

    if not text:
        base = f"{action or 'Assemble'} {default_part}"
        if filled_tools:
            base += f" using {default_tool}"
        return base

    for i in range(20):
        token = f"{{a.{i + 1}}}"
        if token in text:
            slot = part_slots[i] if i < len(part_slots) else None
            val = format_slot(slot) if slot else None
            if not val:
                val = f"[Empty Slot {token}]" if slot else f"[Slot {token} not found]"
            text = text.replace(token, val)

    for i in range(20):
        token = f"{{t.{i + 1}}}"
        if token in text:
            slot = tool_slots[i] if i < len(tool_slots) else None
            val = format_slot(slot) if slot else None
            if not val:
                val = f"[Empty Slot {token}]" if slot else f"[Slot {token} not found]"
            text = text.replace(token, val)

    text = text.replace("{child}", default_part)
    text = text.replace("{a}", default_part)
    text = text.replace("{qty}", "1")
    text = text.replace("{tool}", default_tool)
    text = text.replace("{t}", default_tool)
    text = text.replace("{action}", action or "Assemble")

    return text

def get_recursive_flat_bom(client, parent_id):
    if not client:
        return []
        
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=2) as executor:
        fut_bom = executor.submit(client._get_all_rows, client.table_bom)
        fut_assembly = executor.submit(client._get_all_rows, client.table_assembly)
        bom_rows = fut_bom.result()
        assembly_rows = fut_assembly.result()

    bom_map = {row["id"]: row for row in bom_rows}
    
    parent_to_children = {}
    for edge in assembly_rows:
        p_link = edge.get("Item")
        c_link = edge.get("Contains")
        if p_link and c_link and isinstance(p_link, list) and isinstance(c_link, list) and len(p_link) > 0 and len(c_link) > 0:
            pid = p_link[0].get("id")
            cid = c_link[0].get("id")
            if pid is not None and cid is not None:
                q = edge.get("Amount of Times")
                qty = float(q) if (q is not None and q != "") else 1.0
                
                if pid not in parent_to_children:
                    parent_to_children[pid] = []
                parent_to_children[pid].append({"child_id": cid, "quantity": qty})

    instruction_rows = client._get_all_rows(client.table_instructions)
    items_with_instructions = set()
    for row in instruction_rows:
        p_link = row.get("Parent Item")
        if p_link and isinstance(p_link, list) and len(p_link) > 0:
            items_with_instructions.add(p_link[0].get("id"))

    required_totals = {}
    
    def traverse(current_id, current_multiplier, visited):
        if current_id in visited:
            return
        rels = parent_to_children.get(current_id, [])
        for rel in rels:
            cid = rel["child_id"]
            qty = rel["quantity"] * current_multiplier
            required_totals[cid] = required_totals.get(cid, 0.0) + qty

            child_part = bom_map.get(cid, {})
            is_blackbox = bool(child_part.get("Blackbox", False))
            has_instructions = cid in items_with_instructions

            if not is_blackbox and not has_instructions:
                traverse(cid, qty, visited | {current_id})

    traverse(parent_id, 1.0, set())

    bom_items = []
    for cid, qty_val in required_totals.items():
        c_item = bom_map.get(cid)
        if c_item:
            try:
                qty = int(qty_val) if float(qty_val).is_integer() else qty_val
            except Exception:
                qty = qty_val
            images = c_item.get("Image", [])
            image_url = images[0].get("url") or "" if (images and isinstance(images, list) and len(images) > 0) else ""
            
            part_no = c_item.get("Part Number") or ""
            rev = c_item.get("Revision") or ""
            desc = c_item.get("Item description") or c_item.get("Description") or ""
            ext_pn = c_item.get("External PN") or ""

            bom_items.append({
                "id": cid,
                "part_number": c_item.get("Full PN") or (
                    f"{part_no} Rev.{rev}" if rev else part_no
                ),
                "pn": part_no,
                "revision": rev,
                "description": desc,
                "ext_pn": ext_pn,
                "quantity": qty,
                "image_url": image_url,
                "image": ""
            })
            
    return bom_items

def make_width_filter(doc):
    def width_filter(image_val, size_str):
        if not image_val:
            return ""
        width_cm = 8.0
        if size_str:
            size_clean = size_str.lower().replace("width:", "").replace("cm", "").strip()
            try:
                width_cm = float(size_clean)
            except ValueError:
                pass
        
        if hasattr(image_val, 'width'):
            image_val.width = Cm(width_cm)
            return image_val
            
        if isinstance(image_val, str) and (image_val.startswith("http://") or image_val.startswith("https://") or image_val.startswith("/")):
            try:
                # Use default fallback for API URL since we don't have client here,
                # but if doc has client or environment we could parse it, or default to localhost.
                real_url = get_real_image_url(image_val, "http://localhost:7070")
                resp = requests.get(real_url, timeout=5)
                if resp.status_code == 200:
                    img_stream = io.BytesIO(resp.content)
                    return InlineImage(doc, img_stream, width=Cm(width_cm))
            except Exception as e:
                logger.error(f"Failed to download image from {image_val}: {e}")
            return ""
            
        if isinstance(image_val, str) and os.path.exists(image_val):
            try:
                return InlineImage(doc, image_val, width=Cm(width_cm))
            except Exception as e:
                logger.error(f"Failed to load local image {image_val}: {e}")
            return ""
            
        return ""
    return width_filter

def preprocess_docx_runs(doc):
    pattern = re.compile(r'\|\s*width:([0-9a-zA-Z\.]+)')
    # Fix runs in paragraphs
    for para in doc.paragraphs:
        for run in para.runs:
            if "|" in run.text and "width:" in run.text:
                run.text = pattern.sub(r"| width('\1')", run.text)
                
    # Fix runs in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if "|" in run.text and "width:" in run.text:
                            run.text = pattern.sub(r"| width('\1')", run.text)

    # Fix runs in headers & footers
    for section in doc.sections:
        for para in section.header.paragraphs:
            for run in para.runs:
                if "|" in run.text and "width:" in run.text:
                    run.text = pattern.sub(r"| width('\1')", run.text)
        for para in section.footer.paragraphs:
            for run in para.runs:
                if "|" in run.text and "width:" in run.text:
                    run.text = pattern.sub(r"| width('\1')", run.text)

def move_row_loops_outside(xml_content):
    pattern = re.compile(
        r"(<w:tr[ >](?:(?!<w:tr[ >]).)*?)"
        r"({%\s*for\s+([^%]+?)\s*%})"
        r"(.*?)"
        r"({%\s*endfor\s*%})"
        r"(.*?</w:tr>)",
        flags=re.DOTALL
    )
    def repl(match):
        tr_start = match.group(1)
        loop_start = match.group(2)
        row_body = match.group(4)
        loop_end = match.group(5)
        row_end = match.group(6)
        return f"{loop_start}{tr_start}{row_body}{row_end}{loop_end}"
        
    return pattern.sub(repl, xml_content)

def render_wi_document(template_path, output_path, item, steps, client=None):
    doc = DocxTemplate(template_path)
    
    # Initialize the underlying docx object so it is loaded and not None
    doc.init_docx()
    
    # Override patch_xml to dynamically move table loops outside row XML tags
    original_patch_xml = doc.patch_xml
    def custom_patch_xml(src_xml):
        patched = original_patch_xml(src_xml)
        return move_row_loops_outside(patched)
    doc.patch_xml = custom_patch_xml
    
    # Pre-process the runs to translate | width:8cm syntax to | width('8cm') Jinja filter syntax
    preprocess_docx_runs(doc)
    
    # Register the custom filter on a Jinja2 Environment
    from jinja2 import Environment
    jinja_env = Environment(undefined=SilentUndefined)
    jinja_env.filters['width'] = make_width_filter(doc)
    
    # Get api_url from client if present
    api_url = client.api_url if client else "http://localhost:7070"

    def download_img(url, default_width=5.0):
        if not url:
            return ""
        try:
            real_url = get_real_image_url(url, api_url)
            resp = requests.get(real_url, timeout=5)
            if resp.status_code == 200:
                return InlineImage(doc, io.BytesIO(resp.content), width=Cm(default_width))
        except Exception as e:
            logger.error(f"Failed to download image {url}: {e}")
        return ""

    # Download item image if present
    item_img = ""
    item_images = item.get("Image", [])
    if item_images and isinstance(item_images, list) and len(item_images) > 0:
        item_img = download_img(item_images[0].get("url"), default_width=8.0)
    elif item.get("local_image_path") and os.path.exists(item["local_image_path"]):
        item_img = InlineImage(doc, item["local_image_path"], width=Cm(8.0))

    # Process steps context
    context_steps = []
    all_tools = {}
    for i, step in enumerate(steps):
        # Evaluate step instruction text
        instruction_text = evaluate_instruction_text(step)
        
        # Prepare parts inside step
        parts_list = []
        for p in step.get("part_slots", []):
            if p.get("id") is not None:
                part_img = download_img(p.get("image_url"), default_width=5.0)
                p_qty = p.get("quantity", 1)
                try:
                    p_qty = int(p_qty) if (p_qty is not None and p_qty != "") else 1
                except (ValueError, TypeError):
                    p_qty = 1
                part_no = p.get("pn", "")
                rev = p.get("revision", "")
                full_pn = f"{part_no} Rev.{rev}" if rev else part_no
                parts_list.append({
                    "id": p.get("id"),
                    "pn": part_no,
                    "part_number": p.get("part_number", ""),
                    "revision": rev,
                    "description": p.get("description", ""),
                    "ext_pn": p.get("ext_pn", ""),
                    "quantity": p_qty,
                    "full_pn": full_pn,
                    "image": part_img
                })

        # Prepare tools inside step
        tools_list = []
        for t in step.get("tool_slots", []):
            if t.get("id") is not None:
                tool_img = download_img(t.get("image_url"), default_width=5.0)
                qty = t.get("quantity", 1)
                try:
                    qty = int(qty) if (qty is not None and qty != "") else 1
                except (ValueError, TypeError):
                    qty = 1
                
                # Expose to unique tools list
                tid = t.get("id")
                if tid is not None:
                    if tid not in all_tools:
                        all_tools[tid] = {
                            "id": tid,
                            "pn": t.get("pn", ""),
                            "part_number": t.get("part_number", ""),
                            "revision": t.get("revision", ""),
                            "description": t.get("description", ""),
                            "ext_pn": t.get("ext_pn", ""),
                            "image_url": t.get("image_url", ""),
                            "image": tool_img,
                            "quantity": 0
                        }
                    all_tools[tid]["quantity"] += qty
                
                part_no = t.get("pn", "")
                rev = t.get("revision", "")
                full_pn = f"{part_no} Rev.{rev}" if rev else part_no
                
                tools_list.append({
                    "id": t.get("id"),
                    "pn": part_no,
                    "part_number": t.get("part_number", ""),
                    "revision": rev,
                    "description": t.get("description", ""),
                    "ext_pn": t.get("ext_pn", ""),
                    "quantity": qty if qty > 1 else "",
                    "full_pn": full_pn,
                    "image": tool_img
                })

        step_ctx = {
            "step_number": i + 1,
            "main_action": step.get("action", ""),
            "instruction_text": instruction_text,
            "parts": parts_list,
            "tools": tools_list
        }
        
        # Download step image if present
        step_photo = step.get("photo", [])
        if step_photo and isinstance(step_photo, list) and len(step_photo) > 0:
            step_ctx["step_image"] = download_img(step_photo[0].get("url"), default_width=8.0)
        elif step.get("local_image_path") and os.path.exists(step["local_image_path"]):
            step_ctx["step_image"] = InlineImage(doc, step["local_image_path"], width=Cm(8.0))
        else:
            step_ctx["step_image"] = ""
            
        context_steps.append(step_ctx)

    unique_tools = []
    for tid, ut in all_tools.items():
        qty = ut["quantity"]
        ut["quantity"] = qty if qty > 1 else ""
        part_no = ut.get("pn", "")
        rev = ut.get("revision", "")
        ut["full_pn"] = f"{part_no} Rev.{rev}" if rev else part_no
        unique_tools.append(ut)

    # Fetch recursive flat BOM respecting Blackbox and instruction set boundaries
    bom_items = []
    if client and item.get("id"):
        raw_bom_items = get_recursive_flat_bom(client, item["id"])
        for c in raw_bom_items:
            child_img = download_img(c.get("image_url"), default_width=5.0)
            c["image"] = child_img
            bom_items.append(c)

    context = {
        "item_pn": item.get("Full PN") or (
            f"{item.get('Part Number')} Rev.{item.get('Revision')}"
            if item.get("Revision") else item.get("Part Number", "")
        ),
        "pn": item.get("Part Number", ""),
        "revision": item.get("Revision", ""),
        "full_pn": item.get("Full PN") or (
            f"{item.get('Part Number')} Rev.{item.get('Revision')}"
            if item.get("Revision") else item.get("Part Number", "")
        ),
        "description": item.get("Item description") or item.get("Description", ""),
        "ext_pn": item.get("External PN", ""),
        "date": datetime.now().strftime("%Y-%m-%d"),
        "item_image": item_img,
        "steps": context_steps,
        "step": GlobalStepProxy(context_steps),
        "unique_tools": unique_tools,
        "bom_items": bom_items
    }
    
    doc.render(context, jinja_env=jinja_env)
    doc.save(output_path)
    return context
