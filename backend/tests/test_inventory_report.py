"""
test_inventory_report.py
Tests the 2-tab Inventory Requirement Report generator and API endpoint.
"""

import io
import openpyxl
import pytest
from unittest.mock import MagicMock
from app.inventory_report import generate_inventory_report


def test_generate_inventory_report_structure():
    # Setup mock client
    mock_client = MagicMock()
    mock_client.table_bom = 508
    mock_client.table_assembly = 701

    # BOM items:
    # 1: Root Assembly "80-00001 Rev.A"
    # 2: Sub-Assembly "55-00001 Rev.A" (contains 4 and 5)
    # 3: Blackbox Sub-Assembly "60-00001 Rev.A" (contains 6, but should not recurse!)
    # 4: Leaf Part "20-00001 Rev.A" ($2.50)
    # 5: Leaf Part "30-00001 Rev.A" ($10.00)
    # 6: Child of Blackbox (should NOT appear in BOM!)
    bom_rows = [
        {"id": 1, "Part Number": "80-00001", "Revision": "A", "Full PN": "80-00001 Rev.A", "Item description": "Main Assembly", "Blackbox": False},
        {"id": 2, "Part Number": "55-00001", "Revision": "A", "Full PN": "55-00001 Rev.A", "Item description": "Sub Assembly", "Blackbox": False, "Price per unit": 50.0},
        {"id": 3, "Part Number": "60-00001", "Revision": "A", "Full PN": "60-00001 Rev.A", "Item description": "Blackbox Module", "Blackbox": True, "Price per unit": 120.0, "Sourced By": {"value": "Purchased by ERA"}},
        {"id": 4, "Part Number": "20-00001", "Revision": "A", "Full PN": "20-00001 Rev.A", "Item description": "Screw M3", "Blackbox": False, "Price per unit": 2.5, "External Part Number": "EXT-SCREW-3"},
        {"id": 5, "Part Number": "30-00001", "Revision": "A", "Full PN": "30-00001 Rev.A", "Item description": "Bracket", "Blackbox": False, "Price per unit": 10.0},
        {"id": 6, "Part Number": "20-00002", "Revision": "A", "Full PN": "20-00002 Rev.A", "Item description": "Internal BB Part", "Blackbox": False}
    ]

    # Assembly relations:
    # 1 -> 2 (qty 2)
    # 1 -> 3 (qty 1) [Blackbox]
    # 2 -> 4 (qty 4)
    # 2 -> 5 (qty 1)
    # 3 -> 6 (qty 10) [Should be ignored because 3 is blackbox]
    assembly_rows = [
        {"id": 101, "Item": [{"id": 1}], "Contains": [{"id": 2}], "Amount of Times": 2},
        {"id": 102, "Item": [{"id": 1}], "Contains": [{"id": 3}], "Amount of Times": 1},
        {"id": 103, "Item": [{"id": 2}], "Contains": [{"id": 4}], "Amount of Times": 4},
        {"id": 104, "Item": [{"id": 2}], "Contains": [{"id": 5}], "Amount of Times": 1},
        {"id": 105, "Item": [{"id": 3}], "Contains": [{"id": 6}], "Amount of Times": 10}
    ]

    mock_client._get_all_rows.side_effect = lambda table_id: bom_rows if table_id == 508 else assembly_rows

    stream = generate_inventory_report(mock_client, item_id=1, target_build_qty=10)
    assert isinstance(stream, io.BytesIO)

    wb = openpyxl.load_workbook(stream, data_only=False)
    assert "Nested BOM Requirements" in wb.sheetnames
    assert "Flat BOM Requirements" in wb.sheetnames

    # Verify Tab 1: Nested BOM
    ws1 = wb["Nested BOM Requirements"]
    assert ws1["A1"].value == "80-00001 Rev.A - Main Assembly"
    assert ws1["B2"].value == 10.0

    # Headers at row 4
    headers1 = [ws1.cell(row=4, column=c).value for c in range(1, 11)]
    assert headers1 == ["Level", "ERA PN & Rev", "External PN", "Description", "Sourced By", "Unit Qty", "Total Qty", "Unit Price", "NRE Cost", "Total Cost"]

    # Verify nested rows:
    # Row 5: Level 1 -> 55-00001 (Unit Qty: 2, Total Qty: =F5*$B$2)
    assert ws1["A5"].value == 1
    assert "55-00001 Rev.A" in ws1["B5"].value
    assert ws1["F5"].value == 2.0
    assert ws1["G5"].value == "=F5*$B$2"

    # Row 6: Level 2 -> 20-00001 (Unit Qty: 4, Total Qty: =F6*G5)
    assert ws1["A6"].value == 2
    assert "20-00001 Rev.A" in ws1["B6"].value
    assert ws1["F6"].value == 4.0
    assert ws1["G6"].value == "=F6*G5"

    # Row 7: Level 2 -> 30-00001 (Unit Qty: 1, Total Qty: =F7*G5)
    assert ws1["A7"].value == 2
    assert "30-00001 Rev.A" in ws1["B7"].value
    assert ws1["G7"].value == "=F7*G5"

    # Row 8: Level 1 -> 60-00001 (Blackbox, Unit Qty: 1, Total Qty: =F8*$B$2)
    assert ws1["A8"].value == 1
    assert "60-00001 Rev.A" in ws1["B8"].value
    assert ws1["G8"].value == "=F8*$B$2"

    # Verify that Item 6 (child of blackbox) was NOT traversed
    for r in range(5, ws1.max_row + 1):
        cell_val = ws1.cell(row=r, column=2).value or ""
        assert "20-00002" not in str(cell_val)

    # Verify Tab 2: Flat BOM
    ws2 = wb["Flat BOM Requirements"]
    assert ws2["B2"].value == "='Nested BOM Requirements'!B2"

    headers2 = [ws2.cell(row=4, column=c).value for c in range(1, 12)]
    assert headers2 == ["ERA PN & Rev", "External PN", "Description", "Sourced By", "Unit Price", "Unit Qty / Assy", "Required Qty", "Current Inventory", "Units to Purchase", "Purchase Cost", "Source URL"]

    # Flat items should be sorted alphabetically by Full PN:
    # 20-00001 Rev.A (Unit Qty: 2 * 4 = 8)
    # 30-00001 Rev.A (Unit Qty: 2 * 1 = 2)
    # 60-00001 Rev.A (Blackbox, Unit Qty: 1)
    flat_pns = [ws2.cell(row=r, column=1).value for r in range(5, 8)]
    assert flat_pns == ["20-00001 Rev.A", "30-00001 Rev.A", "60-00001 Rev.A"]

    assert ws2["F5"].value == 8.0  # 20-00001: 2 * 4
    assert ws2["G5"].value == "=F5*$B$2"
    assert ws2["H5"].value == 0
    assert ws2["I5"].value == "=MAX(0, G5-H5)"
    assert ws2["J5"].value == "=IF(ISNUMBER(E5), I5*E5, 0)"

    assert ws2["F6"].value == 2.0  # 30-00001: 2 * 1
    assert ws2["F7"].value == 1.0  # 60-00001: 1


from unittest.mock import patch

@patch('app.main.BaserowClient')
def test_inventory_report_api_endpoint(mock_baserow_client, monkeypatch):
    mock_instance = mock_baserow_client.return_value
    mock_instance.get_item.return_value = {
        "id": 10,
        "Part Number": "55-00007",
        "Revision": "A",
        "Full PN": "55-00007 Rev.A",
        "Item description": "Nova Power Supply"
    }

    # Mock generate_inventory_report
    fake_stream = io.BytesIO(b"PK\x03\x04fake_xlsx_content")
    monkeypatch.setattr("app.inventory_report.generate_inventory_report", lambda c, item_id, target_build_qty: fake_stream)

    from app.main import create_app
    app = create_app()

    with app.test_client() as client:
        resp = client.get("/api/bom/items/10/inventory-report?build_qty=5")
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.content_type
        assert "55-00007 Rev.A - Inventory Requirement Report.xlsx" in resp.headers.get("Content-Disposition", "")


def test_inventory_report_lot_size_price():
    mock_client = MagicMock()
    mock_client.table_bom = 508
    mock_client.table_assembly = 701

    bom_rows = [
        {"id": 1, "Part Number": "80-00001", "Revision": "A", "Full PN": "80-00001 Rev.A", "Item description": "Main Assembly", "Blackbox": False},
        {"id": 2, "Part Number": "20-00001", "Revision": "A", "Full PN": "20-00001 Rev.A", "Item description": "Resistor", "Blackbox": False, "Price per unit": 21.0, "Lot Size": 100.0}
    ]
    assembly_rows = [
        {"id": 101, "Item": [{"id": 1}], "Contains": [{"id": 2}], "Amount of Times": 10}
    ]
    mock_client._get_all_rows.side_effect = lambda table_id: bom_rows if table_id == 508 else assembly_rows

    stream = generate_inventory_report(mock_client, item_id=1, target_build_qty=1)
    wb = openpyxl.load_workbook(stream, data_only=False)
    ws1 = wb["Nested BOM Requirements"]
    # Unit price at row 5 Col H should be 21.0 / 100 = 0.21
    assert ws1["H5"].value == pytest.approx(0.21)


def test_inventory_report_purchase_kit():
    mock_client = MagicMock()
    mock_client.table_bom = 508
    mock_client.table_assembly = 701

    # Root 1 contains Kit 2 (qty 2) and Standalone 5 (qty 3)
    # Kit 2 contains Child 3 (qty 4) and Child 4 (qty 1)
    bom_rows = [
        {"id": 1, "Part Number": "80-00001", "Revision": "A", "Full PN": "80-00001 Rev.A", "Item description": "Main Assembly"},
        {"id": 2, "Part Number": "50-00001", "Revision": "A", "Full PN": "50-00001 Rev.A", "Item description": "Cable Kit", "Purchase Kit": True, "Price per unit": 45.0, "Source URL": "https://supplier.com/kit"},
        {"id": 3, "Part Number": "20-00001", "Revision": "A", "Full PN": "20-00001 Rev.A", "Item description": "Screw M3", "Price per unit": 2.0},
        {"id": 4, "Part Number": "30-00001", "Revision": "A", "Full PN": "30-00001 Rev.A", "Item description": "Bracket", "Price per unit": 10.0},
        {"id": 5, "Part Number": "40-00001", "Revision": "A", "Full PN": "40-00001 Rev.A", "Item description": "Standalone Diode", "Price per unit": 1.5}
    ]
    assembly_rows = [
        {"id": 101, "Item": [{"id": 1}], "Contains": [{"id": 2}], "Amount of Times": 2},
        {"id": 102, "Item": [{"id": 1}], "Contains": [{"id": 5}], "Amount of Times": 3},
        {"id": 103, "Item": [{"id": 2}], "Contains": [{"id": 3}], "Amount of Times": 4},
        {"id": 104, "Item": [{"id": 2}], "Contains": [{"id": 4}], "Amount of Times": 1}
    ]
    mock_client._get_all_rows.side_effect = lambda table_id: bom_rows if table_id == 508 else assembly_rows

    stream = generate_inventory_report(mock_client, item_id=1, target_build_qty=5)
    wb = openpyxl.load_workbook(stream, data_only=False)
    ws2 = wb["Flat BOM Requirements"]

    # Row 5: Kit Header
    assert ws2["A5"].value == "KIT: 50-00001 Rev.A"
    assert ws2["E5"].value == 45.0  # Kit Price
    assert ws2["F5"].value == 2.0   # Kit Unit Qty
    assert ws2["G5"].value == "=F5*$B$2"
    assert ws2["I5"].value == "=MAX(0, G5-H5)"
    assert ws2["J5"].value == "=IF(ISNUMBER(E5), I5*E5, 0)"
    assert ws2["K5"].value == "https://supplier.com/kit"

    # Row 6 & 7: Kit children sorted alphabetically (20-00001 then 30-00001)
    assert "20-00001 Rev.A" in ws2["A6"].value
    assert ws2["E6"].value == 0.0   # Price is 0 for kit child
    assert ws2["F6"].value == 8.0   # 2 * 4 = 8
    assert ws2["I6"].value == "=(F6/F5)*I5"  # Linked to kit purchase count
    assert ws2["J6"].value == 0.0

    assert "30-00001 Rev.A" in ws2["A7"].value
    assert ws2["E7"].value == 0.0
    assert ws2["F7"].value == 2.0   # 2 * 1 = 2
    assert ws2["I7"].value == "=(F7/F5)*I5"

    # Row 8: Standalone Section Header
    assert ws2["A8"].value == "STANDALONE / INDIVIDUAL COMPONENTS"

    # Row 9: Standalone Item (40-00001)
    assert ws2["A9"].value == "40-00001 Rev.A"
    assert ws2["E9"].value == 1.5
    assert ws2["F9"].value == 3.0
    assert ws2["G9"].value == "=F9*$B$2"
    assert ws2["I9"].value == "=MAX(0, G9-H9)"
    assert ws2["J9"].value == "=IF(ISNUMBER(E9), I9*E9, 0)"


import pytest
from unittest.mock import MagicMock
from app.inventory_report import generate_inventory_report
import openpyxl

def test_inventory_report_virtual_purchase_kit():
    client = MagicMock()
    
    client.table_bom = 1
    client.table_assembly = 2
    
    # Mock UOMs
    client.get_uoms.return_value = [
        {"id": 1, "Multiplier to Base": "1.0"},
    ]
    
    # Mock BOM items
    # 1: Top Level Assembly
    # 2: Sub-assembly (NOT a kit)
    # 3: Part A
    # 4: Part B
    # 5: Purchase Kit (contains A and B)
    client._get_all_rows.side_effect = lambda table: {
        1: [
            {"id": 1, "Full PN": "TOP", "Item description": "Top Level", "Purchase Kit": False},
            {"id": 2, "Full PN": "SUB", "Item description": "Sub Assembly", "Purchase Kit": False},
            {"id": 3, "Full PN": "PART_A", "Item description": "Part A", "Purchase Kit": False},
            {"id": 4, "Full PN": "PART_B", "Item description": "Part B", "Purchase Kit": False},
            {"id": 5, "Full PN": "KIT", "Item description": "The Kit", "Purchase Kit": True},
        ],
        2: [
            # Top -> Sub (1 qty)
            {"id": 100, "Item": [{"id": 1}], "Contains": [{"id": 2}], "Amount of Times": 1},
            # Sub -> Part A (2 qty)
            {"id": 101, "Item": [{"id": 2}], "Contains": [{"id": 3}], "Amount of Times": 2},
            # Sub -> Part B (1 qty)
            {"id": 102, "Item": [{"id": 2}], "Contains": [{"id": 4}], "Amount of Times": 1},
            # Kit -> Part A (1 qty)
            {"id": 103, "Item": [{"id": 5}], "Contains": [{"id": 3}], "Amount of Times": 1},
            # Kit -> Part B (1 qty)
            {"id": 104, "Item": [{"id": 5}], "Contains": [{"id": 4}], "Amount of Times": 1},
        ]
    }[table]

    wb_bytes = generate_inventory_report(client, item_id=1, target_build_qty=1.0)
    wb = openpyxl.load_workbook(wb_bytes)
    
    ws1 = wb["Nested BOM Requirements"]
    ws2 = wb["Flat BOM Requirements"]
    
    # Check Flat BOM
    # The kit should be pulled, and we should need 2 kits (to cover 2 Part A)
    # The kit row should be present
    kit_row = None
    part_a_row = None
    part_b_row = None
    
    for row in ws2.iter_rows(min_row=5, values_only=True):
        if row[0]:
            if "KIT: KIT" in str(row[0]):
                kit_row = row
            elif "PART_A" in str(row[0]):
                part_a_row = row
            elif "PART_B" in str(row[0]):
                part_b_row = row

    assert kit_row is not None, "Virtual kit was not pulled into the report"
    assert part_a_row is not None
    assert part_b_row is not None
    
    # Part A is inside the kit. It should show the amount provided by the kit.
    # Kit provides 1 Part A. We buy 2 kits, so Part A provided is 2.
    assert part_a_row[5] == 2.0
    
    # Part B is inside the kit. Kit provides 1 Part B. We buy 2 kits, so Part B provided is 2.
    # Even though we only needed 1, the report shows 2 provided.
    assert part_b_row[5] == 2.0


def test_inventory_report_zero_measurement():
    client = MagicMock()
    client.table_bom = 1
    client.table_assembly = 2
    client.get_uoms.return_value = []

    client._get_all_rows.side_effect = lambda table: {
        1: [
            {"id": 1, "Full PN": "TOP", "Item description": "Top Assembly", "Purchase Kit": False},
            {"id": 2, "Full PN": "DISCRETE_PART", "Item description": "Footswitch", "Purchase Kit": False, "Price per unit": 10.0},
            {"id": 3, "Full PN": "WIRE", "Item description": "Cable Wire", "Purchase Kit": False, "Price per unit": 2.0},
        ],
        2: [
            # Discrete item with Measurement="0.000" (should NOT become 0 qty!)
            {"id": 101, "Item": [{"id": 1}], "Contains": [{"id": 2}], "Amount of Times": 1, "Measurement": "0.000"},
            # Wire with Measurement="250.0" and Amount of Times=2 (should be 2 * 250 = 500)
            {"id": 102, "Item": [{"id": 1}], "Contains": [{"id": 3}], "Amount of Times": 2, "Measurement": "250.0"},
        ]
    }[table]

    wb_bytes = generate_inventory_report(client, item_id=1, target_build_qty=1.0)
    wb = openpyxl.load_workbook(wb_bytes)
    ws1 = wb["Nested BOM Requirements"]
    ws2 = wb["Flat BOM Requirements"]

    # Tab 1: Discrete part Unit Qty must be 1.0 (not 0.0)
    assert ws1["F5"].value == 1.0
    assert ws1["F6"].value == 500.0

    # Tab 2: Discrete part Unit Qty / Assy must be 1.0 (not 0.0)
    rows_tab2 = {row[0]: row for row in ws2.iter_rows(min_row=5, values_only=True) if row[0]}
    assert "DISCRETE_PART" in rows_tab2
    assert rows_tab2["DISCRETE_PART"][5] == 1.0
    assert "WIRE" in rows_tab2
    assert rows_tab2["WIRE"][5] == 500.0
